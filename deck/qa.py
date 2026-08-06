"""Geometry QA: bounds + text-overflow heuristics for the Rivora deck."""
import math
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5

# average char width as a fraction of font size (points)
CHAR_W = {"Courier New": 0.62, "Arial": 0.50}
LINE_H = 1.22  # line height factor


def est_lines(text, font_size, font_face, width_in, char_spacing=0):
    if not text:
        return 0
    cw = CHAR_W.get(font_face, 0.52) * font_size + char_spacing * 0.35
    cpl = max(1, int(width_in * 72 / cw))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, math.ceil(len(para) / cpl))
    return lines


prs = Presentation("Rivora-Deck.pptx")
issues = []
for idx, slide in enumerate(prs.slides, 1):
    for shp in slide.shapes:
        try:
            x, y = shp.left / EMU_IN, shp.top / EMU_IN
            w, h = shp.width / EMU_IN, shp.height / EMU_IN
        except TypeError:
            continue
        name = shp.shape_type
        if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
            issues.append(f"S{idx}: OUT OF BOUNDS {name} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")
        if shp.has_text_frame:
            tf = shp.text_frame
            total_lines = 0
            max_font = 0
            for para in tf.paragraphs:
                text = "".join(r.text for r in para.runs)
                if not text.strip():
                    total_lines += 1 if text == "" and len(tf.paragraphs) > 1 else 0
                    continue
                sizes = [r.font.size.pt for r in para.runs if r.font.size]
                faces = [r.font.name for r in para.runs if r.font.name]
                fs = max(sizes) if sizes else 12
                face = faces[0] if faces else "Arial"
                max_font = max(max_font, fs)
                total_lines += est_lines(text, fs, face, w)
            if max_font and total_lines:
                need = total_lines * max_font * LINE_H / 72
                if need > h * 1.12 and need - h > 0.12:
                    snippet = tf.paragraphs[0].runs[0].text[:40] if tf.paragraphs and tf.paragraphs[0].runs else "?"
                    issues.append(
                        f"S{idx}: TEXT MAY OVERFLOW '{snippet}' box {w:.2f}x{h:.2f}in needs ~{need:.2f}in ({total_lines} lines @ {max_font:.0f}pt)")

if issues:
    print(f"{len(issues)} potential issues:")
    for i in issues:
        print(" -", i)
else:
    print("No geometry issues found.")
